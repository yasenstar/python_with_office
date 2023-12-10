from docx import Document
doc = Document()

table = doc.add_table(5,5)

c1 = table.cell(1,1)
c2 = table.cell(3,3)
# c2.merge(c1)
c1.merge(c2)

doc.save("./7-10-6.docx")