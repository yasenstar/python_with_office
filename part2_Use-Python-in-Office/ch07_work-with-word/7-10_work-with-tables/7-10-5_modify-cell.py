from docx import Document
doc = Document()

table = doc.add_table(3,2)

cell = table.cell(2, 0)

cell.text = "Python"
print(cell.text)

doc.save("./7-10-5.docx")