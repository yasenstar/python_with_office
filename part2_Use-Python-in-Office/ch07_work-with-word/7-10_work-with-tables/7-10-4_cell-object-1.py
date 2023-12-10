from docx import Document
doc = Document()

table = doc.add_table(5,7)

cell1 = table.cell(1,0)
cell2 = table.cell(2,0)

print(type(cell1))

cell1.text = "Mine"
cell2.text = "Author"

doc.save("./7-10-4.docx")